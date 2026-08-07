"""Gera o arquivo examples/Modelo_Importacao_Atos.xlsx.

Executar com: python tools/gerar_modelo_importacao.py
"""

import os

import openpyxl
from openpyxl.styles import Font

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DESTINO = os.path.join(BASE_DIR, "examples", "Modelo_Importacao_Atos.xlsx")


def gerar():
    wb = openpyxl.Workbook()

    dados = wb.active
    dados.title = "Dados"
    dados.append(["Nome", "Grupo", "ContaOrigem"])
    for cell in dados[1]:
        cell.font = Font(bold=True)
    dados.append(["RECONHECIMENTO DE FIRMA", "TABELIONATO DE NOTAS", 1])
    dados.append(["AUTENTICAÇÃO", "TABELIONATO DE NOTAS", 2])
    dados.column_dimensions["A"].width = 40
    dados.column_dimensions["B"].width = 25
    dados.column_dimensions["C"].width = 14

    instrucoes = wb.create_sheet("Instruções")
    linhas = [
        "Instruções de preenchimento",
        "",
        "Nome: recebe o nome da conta (ato) a ser cadastrada.",
        "Grupo: recebe o grupo contábil ao qual a conta pertence.",
        "ContaOrigem: recebe o código da conta utilizado nas tabelas de interface.",
        "",
        "Cada conta deve possuir seu próprio código em ContaOrigem.",
        "Códigos não podem ficar vazios.",
        "Códigos duplicados devem ser conferidos antes da importação — o aplicativo bloqueia",
        "a geração dos scripts enquanto houver códigos duplicados ou ausentes.",
        "",
        "Não inclua outras colunas além de Nome, Grupo e ContaOrigem.",
    ]
    for i, linha in enumerate(linhas, start=1):
        instrucoes.cell(row=i, column=1, value=linha)
        if i == 1:
            instrucoes.cell(row=i, column=1).font = Font(bold=True, size=13)
    instrucoes.column_dimensions["A"].width = 100

    os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
    wb.save(DESTINO)
    print(f"Gerado: {DESTINO}")


if __name__ == "__main__":
    gerar()
