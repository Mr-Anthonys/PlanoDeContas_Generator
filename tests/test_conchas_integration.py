"""Teste de integração com o arquivo real 'Plano de Contas Receitas - Conchas.xlsx'.

Este teste extrai diretamente da planilha original (aba "Script Receitas") os
31 pares Nome/GrupoContábil (linhas 4-34) e os 31 códigos ContaOrigem
(linhas 88-118, mesma posição relativa), monta uma planilha de importação
equivalente ao novo formato (Nome, Grupo, ContaOrigem) e confirma que o
aplicativo:

  * importa a mesma quantidade de contas (31);
  * preserva o código de cada conta, sem inversão ou reordenação;
  * corrige a inconsistência de grupos do arquivo original (PROTESTO nunca era
    criado; REGISTRO CIVIL era criado sem uso) gerando apenas os grupos
    realmente utilizados: TABELIONATO DE NOTAS e PROTESTO;
  * não gera nenhum código automaticamente e não consulta o banco pelo nome.
"""

import os

import openpyxl
import pytest

from src.models.account import GenerationSettings
from src.services import excel_reader, sql_generator, validation_service
from src.services.reference_loader import load_reference_data

CONCHAS_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "Plano de Contas Receitas - Conchas.xlsx",
)

pytestmark = pytest.mark.skipif(
    not os.path.exists(CONCHAS_PATH),
    reason="Arquivo de referência 'Plano de Contas Receitas - Conchas.xlsx' não encontrado.",
)


def _extrair_contas_do_conchas():
    wb = openpyxl.load_workbook(CONCHAS_PATH, data_only=True, read_only=True)
    sheet = wb["Script Receitas"]

    nomes_grupos = []
    for row in range(4, 35):
        nome = sheet.cell(row=row, column=1).value
        grupo = sheet.cell(row=row, column=2).value
        if nome is None or str(nome).strip() == "":
            continue
        nomes_grupos.append((str(nome).strip(), str(grupo).strip()))

    codigos = []
    for row in range(88, 119):
        codigo = sheet.cell(row=row, column=2).value
        if codigo is None or str(codigo).strip() == "":
            continue
        codigos.append(str(int(codigo)) if isinstance(codigo, float) else str(codigo).strip())

    wb.close()
    assert len(nomes_grupos) == len(codigos), "Blocos de Contas e InterfaceContas devem ter o mesmo tamanho no Conchas"
    return [(nome, grupo, codigo) for (nome, grupo), codigo in zip(nomes_grupos, codigos)]


@pytest.fixture(scope="module")
def conchas_rows():
    return _extrair_contas_do_conchas()


@pytest.fixture
def planilha_equivalente(tmp_path, conchas_rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nome", "Grupo", "ContaOrigem"])
    for nome, grupo, codigo in conchas_rows:
        ws.append([nome, grupo, codigo])
    caminho = tmp_path / "conchas_equivalente.xlsx"
    wb.save(caminho)
    return str(caminho)


def test_quantidade_de_contas_igual_ao_original(conchas_rows):
    assert len(conchas_rows) == 31


def test_importacao_preserva_codigo_e_ordem_sem_inversao(planilha_equivalente, conchas_rows):
    raw_rows = excel_reader.read_accounts(planilha_equivalente)
    accounts = validation_service.validate_accounts(raw_rows)
    assert len(accounts) == 31

    for account, (nome_original, _grupo, codigo_original) in zip(accounts, conchas_rows):
        assert account.nome == nome_original
        assert account.conta_origem == codigo_original

    codigos_gerados = [a.conta_origem for a in accounts]
    assert codigos_gerados == [c for _, _, c in conchas_rows]


def test_nenhuma_conta_com_erro_de_validacao(planilha_equivalente):
    raw_rows = excel_reader.read_accounts(planilha_equivalente)
    accounts = validation_service.validate_accounts(raw_rows)
    invalidas = [a for a in accounts if not a.valido]
    assert invalidas == [], f"Contas inválidas inesperadas: {[(a.linha, a.situacao) for a in invalidas]}"


def test_grupos_corrigidos_tabelionato_e_protesto(planilha_equivalente):
    """No arquivo original, os grupos gerados eram TABELIONATO DE NOTAS e
    REGISTRO CIVIL (este último nunca usado pelas contas, e PROTESTO nunca
    era criado apesar de usado por 16 das 31 contas). O aplicativo deve
    corrigir isso gerando apenas os grupos realmente utilizados."""
    raw_rows = excel_reader.read_accounts(planilha_equivalente)
    accounts = validation_service.validate_accounts(raw_rows)

    grupo_portal = sql_generator.generate_grupo_portal(accounts)
    grupo_contabil = sql_generator.generate_grupo_contabil(accounts)

    assert grupo_portal.quantidade == 2
    assert grupo_contabil.quantidade == 2
    assert "TABELIONATO DE NOTAS" in grupo_portal.sql
    assert "PROTESTO" in grupo_portal.sql
    assert "REGISTRO CIVIL" not in grupo_portal.sql
    assert "REGISTRO CIVIL" not in grupo_contabil.sql


def test_interface_contas_e_historico_usam_o_mesmo_codigo_da_linha(planilha_equivalente):
    raw_rows = excel_reader.read_accounts(planilha_equivalente)
    accounts = validation_service.validate_accounts(raw_rows)
    settings = GenerationSettings(
        regime="Interino", tipo_conta="Recebimento", tipo_arq="Emolumentos",
        interface_comum="Emolumentos", interface_forma_pgto="Emolumentos",
        qtd_marcadores=1, historico_texto="EMOLUMENTOS RECEBIDOS - QTD DE ATOS @", historico1="Qtd", historico2="",
    )

    interface_contas = sql_generator.generate_interface_contas(accounts, settings)
    interface_historico = sql_generator.generate_interface_historico(accounts, settings)

    assert interface_contas.quantidade == 31
    assert interface_historico.quantidade == 31

    for account in accounts:
        assert f"N'{account.conta_origem}'" in interface_contas.sql
        assert f"N'{account.conta_origem}'" in interface_historico.sql

    # Nenhum código é gerado automaticamente nem obtido por consulta pelo nome.
    assert "SELECT Codigo" not in interface_contas.sql
    assert "IDENTITY" not in interface_contas.sql


def test_script_completo_gera_sem_erros(planilha_equivalente):
    raw_rows = excel_reader.read_accounts(planilha_equivalente)
    accounts = validation_service.validate_accounts(raw_rows)
    settings = GenerationSettings(
        regime="Interino", tipo_conta="Recebimento", tipo_arq="Emolumentos",
        interface_comum="Emolumentos", interface_forma_pgto="Emolumentos",
        qtd_marcadores=1, historico_texto="EMOLUMENTOS RECEBIDOS - QTD DE ATOS @", historico1="Qtd", historico2="",
    )
    reference = load_reference_data()
    erros = validation_service.validate_settings(settings, reference)
    assert erros == []

    script = sql_generator.generate_full_script(accounts, settings, reference)
    assert script.count("INSERT INTO") >= 31 + 31 + 2 + 2 + 3  # Contas + InterfaceContas/Historico + grupos + interfaces
